"""Rebuild the paper from preserved text and public aggregate figures. No model calls."""
from pathlib import Path
import json, re
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SECTIONS = json.loads((HERE / 'paper-content.json').read_text())
FIGURES = ROOT / 'results/final-20261030-1203/figures'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(20), Mm(20)
sec.left_margin, sec.right_margin = Mm(21), Mm(21)
sec.header_distance, sec.footer_distance = Mm(9), Mm(9)
WIDTH = 168 / 25.4
for name, size, bold in [('Normal',11,False),('Title',19,True),('Subtitle',10.5,False),
                         ('Heading 1',12.5,True),('Heading 2',11.2,True),
                         ('Heading 3',11,True),('Caption',9.5,False)]:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(0,0,0)
    fonts=st.element.get_or_add_rPr().find(qn('w:rFonts'))
    for attr in list(fonts.attrib): del fonts.attrib[attr]
    for key in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+key),'Times New Roman')
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.06
    st.paragraph_format.widow_control = True
doc.styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1','Heading 2','Heading 3']:
    st = doc.styles[name]
    st.paragraph_format.space_before = Pt(11 if name=='Heading 1' else 8)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.keep_with_next = True
doc.styles['Caption'].paragraph_format.line_spacing = 1.05
doc.core_properties.title = 'Hyperparameter Optimization of a Fixed Pi Agent Harness'
doc.core_properties.subject = 'EP6203 academic paper with validation and partial heldout evidence'
doc.core_properties.author = 'EP6203 Project Team'
doc.core_properties.keywords = 'hyperparameter optimization, agent harness, IGBS, SkillsBench'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)

def clean(text):
    def refs(match):
        return '[' + re.sub(r'S([1-6])',lambda x:str(int(x[1])+3),match[0][1:-1]) + ']'
    text = re.sub(r'\[[^\]]*S[1-6][^\]]*\]', refs, text)
    text = text.replace('Section 8','Section 6').replace('Section 3','Section 4')
    text=text.replace('report-curve-data.csv','aggregate-tables/common-budget-curves.csv')
    text=text.replace('Source hashes, attempt evidence and independent audit records are retained with this report.',
        'Public source hashes and aggregate audit records accompany this paper; raw attempt records remain in the local evidence archive.')
    return text.replace('Hong Kong time','UTC+8').replace('Hong Kong','UTC+8')

def p(text, style=None, center=False):
    para = doc.add_paragraph(clean(text), style)
    if center: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para

def h(text, level=1):
    para=doc.add_heading(text, level)
    if text in ('References','Appendix A Requirements and Reproducibility'):
        para.paragraph_format.page_break_before=True
    return para

CAPTIONS = {
0:'Validation coverage and measurement completeness by method.',
1:'Exposed configuration parameters and declared domains.',
2:'The eight predeclared validation task identifiers.',
3:'Configuration generation methods used in the comparison.',
4:'Outcome treatment under the revised operational policy.',
5:'Frozen validation selections. Valid-only fractions are descriptive and do not replace the fixed-denominator selection metric.',
6:'Partial held-out coverage. Ended-subset fractions are not final eight-task pass@1.',
7:'Outcomes on the three held-out tasks that ended for each method. EXC denotes an accounting exception.',
8:'Recorded stage costs, unknown-request reserves and allocation caps in USD.',
9:'Selected-configuration runtime cost on three ended held-out tasks per method.',
10:'Exploratory reasoning-level comparison at fixed compaction settings.',
11:'Representative validation outcome differences for the frozen configurations.',
12:'Course requirement coverage and remaining evidence gaps.'}
WEIGHTS = {0:[1,1.3,1.3,1.3,1.3],1:[2.3,2.1,2.2],2:[1,1],3:[1,5.6],4:[2.2,4.4],
5:[1.5,.9,.95,.95,1.15,1.15],6:[1,1,.7,1,1.5,1.4],7:[2.9,.9,.9,.9,.9],
8:[1.2,1.4,1.2,1.6,.7],9:[1,1.8,.9,1,1.7],10:[1.1,.9,1.1,1.4,1.7],11:[2.8,1,1.4,1.4],12:[2,2.3,2.3]}
table_number = 0

def table(block):
    global table_number
    table_number += 1
    n = block['id']
    cap = p(f'Table {table_number}. {CAPTIONS[n]}','Caption')
    cap.paragraph_format.keep_with_next = True
    rows = block['rows']
    widths = [WIDTH*x/sum(WEIGHTS[n]) for x in WEIGHTS[n]]
    t = doc.add_table(rows=0,cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for i,values in enumerate(rows):
        row=t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for side,size in [('top',65),('bottom',65),('left',80),('right',80)]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),str(size));el.set(qn('w:type'),'dxa');mar.append(el)
            pr.append(mar)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F0');pr.append(shade)
            para=cell.paragraphs[0]
            para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.03
            para.paragraph_format.keep_with_next=i<len(rows)-1
            para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or n in (1,2,3,4,12) else WD_ALIGN_PARAGRAPH.CENTER
            run=para.add_run(clean(str(value)));run.font.size=Pt(9.5);run.bold=i==0
    gap=p('');gap.paragraph_format.space_after=Pt(0);gap.paragraph_format.line_spacing=Pt(3)
    return t

def emit(section, level=2, omit_headings=(), stop_before=None):
    for b in SECTIONS[section]:
        if b['kind']=='table': table(b)
        elif b['kind']=='figure':
            para=p('');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next=True
            shape=para.add_run().add_picture(str(FIGURES/(b['name']+'.png')),width=Inches(WIDTH))
            shape._inline.docPr.set('descr',CAPTIONS[5] if 'frozen' in b['name'] else 'Validation incumbents on shared supported cost intervals.')
        else:
            text=b['text'];style=b['style']
            if stop_before and text==stop_before: break
            if style.startswith('Heading'):
                if text not in omit_headings: h(text,level)
            else:
                p(text,'Caption' if style=='Caption' else None)

p('Hyperparameter Optimization of a Fixed Pi Agent Harness','Title',True)
p('EP6203 Project Team','Subtitle',True)
p('Nanyang Technological University  |  30 October 2026','Subtitle',True)
h('Abstract')
p('We study hyperparameter optimization over supported configuration interfaces of a fixed Pi Agent harness on a reproducible SkillsBench 1.1 subset. A 27-point space combines reasoning level, context reserve and recent-token retention. We implement Interaction-Guided Block Search (IGBS) and compare it with an untuned default, random search and standard SMAC Bayesian optimization under equal search allocation caps. The recorded campaign ended with 139 validation attempts and 12 held-out attempts; 136 of these 151 observations have valid measurement. All four configurations were frozen before testing. IGBS selected a configuration with 5/8 validation successes, but BO selected identical settings and observed 3/8, so the difference does not establish algorithm superiority. IGBS reached only initialization screening. Three of eight held-out tasks ended per method, leaving final pass@1 and the full generalization gap unavailable. Ended observations record USD 20.4248448 in usage-based estimates and USD 17 in unknown-request reserves. The study provides an auditable evaluation pipeline, exploratory parameter and cost analyses, and verified failure cases, while identifying measurement completeness and evaluation noise as limits on the conclusions.')
kp=p('Keywords: hyperparameter optimization; agent harness; Bayesian optimization; interaction analysis; SkillsBench.')
kp.runs[0].italic=True;kp.runs[0].font.size=Pt(10)

h('1 Introduction')
p('Agent harnesses expose runtime controls that can change task completion and resource use without changing the underlying model. This project treats their supported configuration interfaces as a finite hyperparameter optimization problem. The harness implementation, model, task environments and verifiers remain fixed while configurations are selected using validation outcomes. The selected configurations are then evaluated on a separate held-out task panel [7].')
p('Our question is whether a small-budget search over reasoning and context-management settings can identify a useful configuration and how its observed improvement relates to one-time search expense and subsequent runtime expense. We develop IGBS, a policy that uses local matched parameter contrasts to choose between pair, single-parameter and global moves. Its implemented adaptive mechanism must be distinguished from the stages actually reached during the recorded run.')
p('The campaign stopped at 12:03:07 UTC+8 on 30 October 2026 after an infrastructure error. We report the completed validation panels, partial independent test observations and accounting exceptions explicitly. The results support a study of configuration behavior and measurement challenges, but do not establish a final optimizer ranking. All claims in this paper refer to this fixed evidence cutoff [4–6].')

h('2 Related Work')
p('Random search is a natural baseline for evaluating adaptive hyperparameter optimization. Bergstra and Bengio [1] show its effectiveness relative to grid and manual search in their learning-system experiments. We use seeded sampling without replacement from the finite 27-point domain; their empirical findings do not imply a performance advantage on our agent tasks.')
p('SMAC3 provides a modular Bayesian optimization framework with surrogate models, acquisition rules and configuration-space support [2]. We use the pinned SMAC implementation as the required standard BO baseline. IGBS instead uses directly observed task-paired contrasts and an interpretable proposal schedule. BO can also model parameter interactions, so the contribution of IGBS is its particular routing policy rather than the general idea of considering interactions.')
p('Hutter, Hoos and Leyton-Brown [3] use functional ANOVA to decompose predicted performance variation into parameter effects and interactions. Our observed local contrasts are more restricted: they compare matched settings on the validation tasks and do not estimate global fANOVA importance. The small sample and incomplete mechanism activation motivate exploratory sensitivity analysis rather than a causal ranking of parameters.')

h('3 Experimental Setup')
emit('1 Experiment design and exposed configuration')
h('4 Optimization Methods')
h('Baselines and Measurement Policy',2)
emit('2 Baselines and measurement policy',3)
h('Interaction Guided Block Search',2)
emit('3 IGBS algorithm and execution status',3)

h('5 Experimental Results')
h('Validation Coverage and Frozen Configurations',2)
p('Table 5 summarizes the observed validation coverage. The stopped run contains 139 ended validation attempts, followed by 12 ended held-out attempts; three historical BO attempts are reused once within these counts. Across both stages, 136 ended attempts have valid measurement and 15 retain accounting exceptions [4, 6].')
table(next(b for b in SECTIONS['Executive findings'] if b['kind']=='table'))
emit('4 Frozen validation configurations',3)
h('Incumbent Curves Under Comparable Budgets',2)
emit('5 Incumbent curves under comparable budgets',3)
h('Partial Heldout Evaluation',2)
emit('6 Partial heldout evaluation',3)
h('Search Cost and Selected Runtime Cost',2)
emit('7 Search cost and selected runtime cost',3)

h('6 Parameter Sensitivity')
emit('8 Parameter sensitivity and runtime mechanisms')
h('7 Failure Analysis')
emit('9 Failure cases and evidence limits',2,stop_before='Infrastructure failure and validity threats')
h('Verified Trajectory Cases',2)
emit('10 Verified trajectory cases',3)

h('8 Discussion and Limitations')
take=False
for b in SECTIONS['9 Failure cases and evidence limits']:
    if b.get('text')=='Infrastructure failure and validity threats':take=True;continue
    if take and b['kind']=='paragraph':p(b['text'])
p('The common-budget curves distinguish known usage-price estimates from commitment including reserves. They cannot establish an equal-invoice-cost ranking while provider charges remain unresolved. Likewise, the three-task held-out prefix cannot replace the predeclared eight-task denominator. The original strict protocol, full held-out pass@1 and generalization gap remain incomplete [4–6].')
p('A further limitation concerns the proposed method itself. IGBS completed only five screening panels and part of a sixth. Its interaction-guided adaptive phase was never executed, so the present evidence tests initialization observations rather than the efficacy of adaptive routing. Repeated measurements, completed test coverage and active-mechanism contrasts are needed before stronger conclusions are justified.')
h('9 Conclusion')
p('We implemented a reproducible configuration-only HPO pipeline for Pi Agent and compared Default, Random, SMAC BO and IGBS on the same declared space. The recorded study provides frozen validation selections, cost curves on common observed support, exploratory sensitivity results and verified failure cases. BO and IGBS selected identical settings yet obtained different outcomes, and IGBS did not reach its adaptive phase. Combined with incomplete held-out coverage and unresolved accounting, these observations preclude a final superiority or generalization claim. The main result is an inspectable partial study and a clear account of the evidence still required.')

h('References')
refs=[
'[1] J. Bergstra and Y. Bengio. Random Search for Hyper-Parameter Optimization. Journal of Machine Learning Research, 13:281–305, 2012. https://www.jmlr.org/papers/v13/bergstra12a.html',
'[2] M. Lindauer, K. Eggensperger, M. Feurer, A. Biedenkapp, D. Deng, C. Benjamins, T. Ruhkopf, R. Sass, and F. Hutter. SMAC3: A Versatile Bayesian Optimization Package for Hyperparameter Optimization. Journal of Machine Learning Research, 23(54):1–9, 2022. https://www.jmlr.org/papers/v23/21-0888.html',
'[3] F. Hutter, H. Hoos, and K. Leyton-Brown. An Efficient Approach for Assessing Hyperparameter Importance. Proceedings of the 31st International Conference on Machine Learning, PMLR 32(1):754–762, 2014. https://proceedings.mlr.press/v32/hutter14.html',
'[4] EP6203 Project Team. Recorded campaign aggregates at 12:03 UTC+8, 30 October 2026. Supplementary repository: results/final-20261030-1203/aggregate-tables/ and figure-input.json. These files report validation, partial held-out coverage, operational outcomes and cost arithmetic.',
'[5] EP6203 Project Team. Experiment configuration and frozen provenance, 30 October 2026. Supplementary repository: results/final-20261030-1203/experiment-config.json and source-provenance.json; selected parameters are in aggregate-tables/frozen-winners.json.',
'[6] EP6203 Project Team. Independent evidence audit, 30 October 2026. Supplementary repository: results/final-20261030-1203/audit.json and aggregate-tables/. The audit records checks of 160 protocol source hashes and 1,267 frozen evidence hashes. Raw private records are retained locally.',
'[7] B. An. Group Project 3 Harness Hyperparameter Optimization. NTU course assignment brief, 2026. Problem description, evaluation and deliverables. The instructor document is cited but not redistributed.',
'[8] EP6203 Project Team. HPO implementation and operational continuation. Supplementary repository: hpo/methods.py, hpo/interaction_search.py, hpo/contrasts.py and scripts/restart_20261029/deadline_run.py. Source-level design is distinguished from observed execution.',
'[9] EP6203 Project Team. Verified trajectory case summaries, 30 October 2026. Supplementary repository: results/final-20261030-1203/trajectory-report-paragraphs.json and aggregate-tables/trajectory-cases.json. Summaries include relative references to locally retained evidence.'
]
for ref in refs:
    para=p(ref)
    para.paragraph_format.left_indent=Inches(.22);para.paragraph_format.first_line_indent=Inches(-.22)
    para.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for run in para.runs:run.font.size=Pt(9.5)

h('Appendix A Requirements and Reproducibility')
table(next(b for b in SECTIONS['11 Course requirements and reproducibility'] if b['kind']=='table'))
p('The submission repository contains the HPO implementation, pinned configuration and dependency files, portable aggregate data, the presentation and this paper in Word and PDF form. The requirement checklist in docs/submission_requirements.md distinguishes implemented software, available evidence and incomplete final evaluation requirements. Packaging does not complete missing experiments.')
p('For an offline package integrity check, run python -B scripts/verify_release_package.py from the repository root. The reproduction guide is docs/reproduction.md. The script results/final-20261030-1203/reproduce_figures.py regenerates the two figures from public aggregate inputs. reports/source/build_paper.py rebuilds the editable paper from paper-content.json and those figures; python-docx is required. These checks and report builds make no model calls.')
p('A native benchmark rerun additionally requires the pinned upstream sources, task environments, local credentials and a separately declared budget. The dated controller is preserved as historical execution code and must not be relaunched against the stopped campaign. Raw trajectories and request logs remain in the original local workspace; public source hashes establish file identity, not availability of missing observations or reconciled provider charges.')

in_appendix=False
for paragraph in doc.paragraphs:
    if paragraph.text=='Appendix A Requirements and Reproducibility': in_appendix=True
    if in_appendix and paragraph.style.name=='Normal':
        paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
for parent in (doc.element,doc.styles.element):
    for node in list(parent.iter(qn('w:pBdr'))):node.getparent().remove(node)
target=ROOT/'reports/EP6203_Paper.docx'
doc.save(target)
print(target)

# Integrate the separately measured synthetic validation after the base paper.
import runpy
runpy.run_path(str(HERE/'integrate_synthetic.py'))
