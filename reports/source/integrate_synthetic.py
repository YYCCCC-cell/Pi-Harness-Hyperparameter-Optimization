from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=Path(__file__).resolve().parents[2]
DEMO=OUT/'results/igbs-synthetic-20261030'
d=Document(OUT/'reports/EP6203_Paper.docx')
for p in d.paragraphs:
 t=p.text
 if t.startswith('We study hyperparameter optimization'):
  p.text=t.replace('IGBS reached only initialization screening.','The real IGBS run reached only initialization screening. A separate controlled synthetic test completed adaptive routing and verified strong-signal weighting and conservative fallback (Appendix B).')
 elif t.startswith('Under this continuation,'):
  p.text=t.replace('The reported results cannot validate interaction-guided routing.','These real-task results cannot validate interaction-guided routing; Appendix B separately tests its implementation on controlled synthetic inputs.')
 elif t.startswith('IGBS’s source supports'):
  p.text=t.replace('were not reached.','were not reached in the real campaign. The separate synthetic validation in Appendix B verifies routing behavior, without establishing real-task efficacy.')
 elif t.startswith('A further limitation concerns'):
  p.text='A further limitation concerns empirical efficacy. IGBS completed only five real-task screening panels and part of a sixth, so its adaptive phase was not executed on SkillsBench. A separate controlled synthetic experiment completed the adaptive cycle and verified weighting, fallback and pair-blind ablation behavior (Appendix B). Artificial groups and a signal deliberately aligned with the screening design do not establish real-task performance gains. Repeated measurements and completed independent test coverage remain necessary.'
 elif t.startswith('We implemented a reproducible configuration-only'):
  p.text=t.replace('IGBS did not reach its adaptive phase.','IGBS did not reach its adaptive phase in the real campaign. The added controlled synthetic study verifies the adaptive implementation: strong interactions receive positive pair weight, while absent or uncertain signals trigger uniform fallback. This evidence supports mechanism correctness, not optimizer superiority.')
# Dedicated appendix avoids mixing synthetic panels into primary outcome tables.
d.add_page_break()
d.add_heading('Appendix B Controlled Synthetic Mechanism Validation',level=1)
d.add_paragraph('We executed the unchanged IGBS implementation on three declared synthetic objectives with a matched pair-blind ablation. Each of the six runs used seed 31, eight artificial task groups and ten unique configuration panels: seven screening proposals, followed by pair, single and global proposals. This produced 60 complete panels and 480 synthetic evaluations. The local algorithm run took 0.262 seconds, excluding plotting and packaging; it made zero model calls and incurred USD 0 in paid API costs.')
d.add_paragraph('The 27-point space uses reasoning levels low/medium/high, reserveTokens 16384/360000/375000, and keepRecentTokens 4000/10000/20000. An AND interaction between high reasoning and reserveTokens 360000 was deliberately aligned with the fixed initial screening design. All eight groups succeed at the target pair in the strong case, none succeed in the flat case, and only two groups succeed in the weak case. These are known-answer checks, not an unbiased benchmark of search efficiency.')
d.add_paragraph('Table 14. Observed first adaptive decision on synthetic inputs. Blind denotes the applied pair weight after the ablation.',style='Caption')
t=d.add_table(rows=1,cols=5);t.style='Table Grid'
for c,v in zip(t.rows[0].cells,['Synthetic case','Contrast','Bootstrap radius','IGBS weight','Blind weight']):c.text=v
for row in [('Strong','1.00','0.00','0.95','0.00'),('Flat','0.00','0.00','0.00','0.00'),('Weak uncertain','0.25','0.25','0.00','0.00')]:
 for c,v in zip(t.add_row().cells,row):c.text=v
for row in t.rows:
 for c in row.cells:
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(3)
   for r in p.runs:r.font.size=Pt(9)
d.add_paragraph('IGBS selected the injected parameter pair with weight 0.95 in the strong case. The pair-blind control removed pair weighting and used uniform selection. Both flat and weak cases produced zero IGBS pair weights and identical ten-configuration trajectories across the two methods. The original thresholds were retained: eight groups, practical threshold 0.05, stability threshold 0.8 and 2,000 bootstrap samples.')
d.add_picture(str(OUT/'reports/figures/mechanism-check-no-watermark.png'),width=Inches(6.35))
d.add_paragraph('Figure 3. Controlled synthetic mechanism check. Artificial groups are not independent empirical replicates; bootstrap radii are algorithm outputs on the constructed inputs, not confidence bounds for real-task improvement.',style='Caption')
d.add_paragraph('All six trace hash chains and completed ask/tell sequences passed verification, and 165 protected source files were unchanged. Network and subprocess execution were blocked. The protocol, per-round configurations, traces and checks are retained under results/igbs-synthetic-20261030. This validates implemented routing and ablation behavior only; it does not demonstrate SkillsBench gains, a compaction effect, generalization, or superiority over BO, Random or pair-blind.')

for p in d.paragraphs:
 if 'All claims in this paper refer to this fixed evidence cutoff' in p.text:
  p.text=p.text.replace('All claims in this paper refer to this fixed evidence cutoff','Real-campaign claims refer to this fixed evidence cutoff')
for table in d.tables:
 for row in table.rows:
  for cell in row.cells:
   if cell.text=='IGBS screening only; adaptive stage untested':
    cell.text='Real run: screening only. Synthetic routing verified in Appendix B.'
    for para in cell.paragraphs:
     for run in para.runs:run.font.size=Pt(9.5)
t=d.tables[-1]
borders=OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
 el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
t._tbl.tblPr.append(borders)
for cell in t.rows[0].cells:
 shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F0');cell._tc.get_or_add_tcPr().append(shade)
 for para in cell.paragraphs:
  for run in para.runs:run.bold=True

d.save(OUT/'reports/EP6203_Paper.docx')
